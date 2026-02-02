from openpyxl import load_workbook
import json
import re

workbook = load_workbook("Параметры.xlsx")['Лист3']

all_param = dict()

for i in range(1, 256):
    all_param[i] = {'Байт': "", "Параметры": []}

for row in workbook.iter_rows(min_row=4, max_row=workbook.max_row, values_only=True):
    type_data = "int"
    if (row[6].find("0 -") != -1 and (row[6].find("1 -")!= -1 or row[6].find("1-") != -1))  or (row[6].find("0 —") != -1 and row[6].find("1 —")) or (row[6].find("0 –") != -1 and row[6].find("1 –")):
    # if ("–" in row[6] or "—" in row[6] or '-' in row[6]) and ("0" in row[6] or "1" in row[6]):
        type_data = "bool"
    elif row[3].find("Float") != -1:
        type_data = "float"
    all_param[int(row[0])]['Байт'] = row[2]
    if 1 not in all_param[int(row[0])]['Параметры']:
        all_param[int(row[0])]['Параметры'] = {1: {"Название": row[7], "Бит": str(row[5]), "Тип данных": type_data}}
    else:
        all_param[int(row[0])]['Параметры'][list(all_param[int(row[0])]['Параметры'].keys())[-1]+1] = {"Название": row[7], "Бит": str(row[5]), "Тип данных": type_data}
with open("Параметры.json", "w") as file:
    file.write(json.dumps(all_param))


print(len(all_param[4]['Параметры']))